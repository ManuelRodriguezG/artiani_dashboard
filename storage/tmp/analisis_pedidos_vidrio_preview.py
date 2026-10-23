import json
from pathlib import Path

from openpyxl import load_workbook

root = Path(r"C:\xampp\htdocs\panel_de_control")
source = Path(r"C:\Users\aleja\Downloads\Pedidos Peceras.xlsx")
output_json = root / "docs" / "erp_produccion_pedidos_vidrio_preview_2026-10-22.json"
output_md = root / "docs" / "erp_produccion_pedidos_vidrio_preview_2026-10-22.md"

workbook = load_workbook(source, read_only=True, data_only=False)
pedidos = []

def text(value):
    return str(value).strip() if value is not None else ""

def number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0

for sheet in workbook.worksheets:
    sheet_name = sheet.title
    if not sheet_name.lower().startswith("pedido") or sheet_name.lower().startswith("copia de"):
        continue
    current = None
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        values = list(row)
        d = text(values[3].value if len(values) > 3 else "")
        if d and "hoja" in d.lower():
            if current and current["piezas"]:
                current["piezas_total"] = sum(item["cantidad"] for item in current["piezas"])
                current["espesores_mm"] = sorted({item["espesor_mm"] for item in current["piezas"]})
                pedidos.append(current)
            current = {
                "hoja_origen": sheet_name,
                "fila_inicio": row_number,
                "titulo": d,
                "estatus_origen": text(values[18].value if len(values) > 18 else ""),
                "hoja_largo_m": 0,
                "hoja_ancho_m": 0,
                "piezas": [],
            }
            continue
        if not current:
            continue
        codigo = text(values[1].value if len(values) > 1 else "")
        tipo = text(values[2].value if len(values) > 2 else "")
        largo = number(values[3].value if len(values) > 3 else "")
        ancho = number(values[4].value if len(values) > 4 else "")
        cantidad = max(1, round(number(values[5].value if len(values) > 5 else "")))
        espesor = number(values[6].value if len(values) > 6 else "")
        if codigo and tipo and largo > 0 and ancho > 0 and espesor > 0:
            current["piezas"].append({
                "fila": row_number,
                "codigo_crudo": codigo,
                "tipo_crudo": tipo,
                "largo_cm": largo,
                "ancho_cm": ancho,
                "cantidad": cantidad,
                "espesor_mm": espesor,
                "cantidad_por_hoja": max(0, round(number(values[17].value if len(values) > 17 else ""))),
            })
        for largo_index, ancho_index in ((11, 12), (20, 21)):
            largo_hoja = number(values[largo_index].value if len(values) > largo_index else "")
            ancho_hoja = number(values[ancho_index].value if len(values) > ancho_index else "")
            if 0 < largo_hoja <= 10 and 0 < ancho_hoja <= 10:
                current["hoja_largo_m"] = max(current["hoja_largo_m"], largo_hoja)
                current["hoja_ancho_m"] = max(current["hoja_ancho_m"], ancho_hoja)
    if current and current["piezas"]:
        current["piezas_total"] = sum(item["cantidad"] for item in current["piezas"])
        current["espesores_mm"] = sorted({item["espesor_mm"] for item in current["piezas"]})
        pedidos.append(current)

preview = {
    "fuente": "Pedidos Peceras.xlsx",
    "fecha_analisis": "2026-10-22",
    "solo_lectura": True,
    "bloques_pedido": len(pedidos),
    "pedidos": pedidos,
}
output_json.write_text(json.dumps(preview, ensure_ascii=False, indent=2), encoding="utf-8")

lines = [
    "# Preview de pedidos historicos de vidrio",
    "",
    "Fuente: Pedidos Peceras.xlsx",
    "",
    "Preview de solo lectura. Conserva cada bloque de hoja por separado y no modifica perfiles ni la base de datos.",
    "",
    f"Bloques de pedido encontrados: **{len(pedidos)}**",
    "",
    "| Hoja | Fila bloque | Espesores | Medida hoja | Piezas | Estado |",
    "|---|---:|---|---|---:|---|",
]
for pedido in pedidos:
    measure = (
        f'{pedido["hoja_largo_m"]} x {pedido["hoja_ancho_m"]} m'
        if pedido["hoja_largo_m"] else "No detectada"
    )
    status = pedido["estatus_origen"] or "Sin estado"
    lines.append(
        f'| {pedido["hoja_origen"]} | {pedido["fila_inicio"]} | '
        f'{", ".join(str(item) for item in pedido["espesores_mm"])} mm | '
        f'{measure} | {pedido["piezas_total"]} | {status} |'
    )
output_md.write_text("\n".join(lines) + "\n", encoding="utf-8")
print(json.dumps({"json": str(output_json), "markdown": str(output_md), "bloques": len(pedidos)}, ensure_ascii=False))
